"""Genera el SQL para cargar terceros y cuentas de la empresa (A&F, migración
20261016000000_ayf_terceros.sql) desde BASE_TERCEROS.xlsx y BASE_CUENTAS.xlsx,
las bases de la app de dispersores.

Uso (desde app-demo/):
    python scripts/generar_seed_terceros.py [carpeta_config] [salida.sql]

Por defecto lee de la carpeta config de la app de dispersores y escribe en
referencias/seed_terceros_ayf.sql (esa carpeta está en .gitignore: el SQL trae
números de cuenta reales y NO debe subirse a git).

El SQL se ejecuta a mano en el SQL Editor de Supabase, DESPUÉS de la migración.
Es repetible: no duplica lo que ya está cargado. Al final muestra conteos y
las revisiones pendientes.
"""

import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

CONFIG_POR_DEFECTO = Path(
    r"G:\Shared drives\3. CONTABILIDAD\CONTABILIDAD\6. CUADRO DE PAGOS\Dispersores\config")
SALIDA_POR_DEFECTO = Path(__file__).resolve().parent.parent / "referencias" / "seed_terceros_ayf.sql"

# Cómo aparece escrito en la base de terceros -> nombre canónico de public.bancos
ALIAS_BANCO = {
    "BANCO DAVIVIENDA": "DAVIVIENDA",
    "BANCO AV VILLAS": "AV VILLAS",
    "NU": "NU BANK",
}
TIPOS_DOC = {"CC", "NIT", "CE", "PPT", "PA", "TI", "RC"}
TIPOS_CUENTA = {"AHORROS", "CORRIENTE"}
LOTE = 400  # filas por INSERT


def texto(v) -> str:
    s = str(v if v is not None else "").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip().upper()


def digitos(v) -> str:
    s = str(v if v is not None else "").replace("\xa0", " ").strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    return re.sub(r"\D", "", s)


def fecha(v) -> str:
    if v is None or str(v).strip() == "":
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()[:10]
    return s if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s) else ""


def q(v) -> str:
    """Literal SQL de texto ('' si vacío -> NULL)."""
    if v is None or v == "":
        return "null"
    return "'" + str(v).replace("'", "''") + "'"


def lotes(filas, n=LOTE):
    for i in range(0, len(filas), n):
        yield filas[i:i + n]


def leer_terceros(ruta, avisos):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["TERCEROS"] if "TERCEROS" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(min_row=2, values_only=True))

    # unifica "Luis Perez" / "LUIS PEREZ" en la grafía más frecuente
    grafias = Counter(str(f[13]).strip() for f in filas if f[13])
    canon = {}
    for g, _ in grafias.most_common():
        canon.setdefault(g.lower(), g)

    terceros = {}   # documento -> dict
    cuentas = {}    # (documento, banco, cuenta) -> dict
    for i, f in enumerate(filas, start=2):
        if not any(f):
            continue
        doc = digitos(f[1])
        if not doc:
            continue
        tipo = texto(f[0]) or "CC"
        if tipo not in TIPOS_DOC:
            avisos.append(f"fila {i}: tipo de documento '{tipo}' no válido ({doc}); se omite.")
            continue
        nombre = texto(f[3])
        if not nombre:
            avisos.append(f"fila {i}: documento {doc} sin nombre; se omite.")
            continue

        t = terceros.get(doc)
        if t is None:
            terceros[doc] = {
                "tipo": tipo, "doc": doc, "dv": digitos(f[2]), "nombre": nombre,
                "nombre_banco": texto(f[4]), "nombres": texto(f[5]),
                "apellidos": texto(f[6]),
            }
        else:
            if t["tipo"] != tipo:
                avisos.append(f"documento {doc}: aparece como {t['tipo']} y como {tipo} (fila {i}).")
            if t["nombre"] != nombre:
                avisos.append(f"documento {doc}: nombres distintos ('{t['nombre']}' / '{nombre}', fila {i}); "
                              f"se conserva el primero.")

        banco = texto(f[7])
        banco = ALIAS_BANCO.get(banco, banco)
        cuenta = digitos(f[9])
        tipo_c = texto(f[8])
        if not cuenta:
            avisos.append(f"fila {i}: {nombre} sin número de cuenta; se carga el tercero sin cuenta.")
            continue
        if tipo_c not in TIPOS_CUENTA:
            avisos.append(f"fila {i}: {nombre} tipo de cuenta '{tipo_c}' no válido; se omite la cuenta.")
            continue
        clave = (doc, banco, cuenta)
        if clave in cuentas:
            avisos.append(f"fila {i}: cuenta repetida de {nombre} en {banco}; se conserva la primera.")
            continue
        estado = texto(f[11]) or "ACTIVO"
        if estado not in ("ACTIVO", "PENDIENTE", "INACTIVO"):
            estado = "PENDIENTE"
        ver_por = canon.get(str(f[13]).strip().lower(), str(f[13] or "").strip()) if f[13] else ""
        cuentas[clave] = {
            "doc": doc, "banco": banco, "tipo": tipo_c, "cuenta": cuenta,
            "etiqueta": str(f[10] or "").strip(), "estado": estado,
            "ver_en": fecha(f[12]), "ver_por": ver_por,
            "obs": str(f[14] or "").strip(),
        }
    return list(terceros.values()), list(cuentas.values())


def leer_cuentas_empresa(ruta, avisos):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["CUENTAS"] if "CUENTAS" in wb.sheetnames else wb.worksheets[0]
    out, sin_cuenta = [], []
    for i, f in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(f):
            continue
        apodo = texto(f[0])
        if not apodo:
            sin_cuenta.append(texto(f[7]))
            continue
        banco = texto(f[2])
        out.append({
            "apodo": apodo, "dim": texto(f[1]), "banco": banco, "tipo": texto(f[3]),
            "cuenta": digitos(f[4]), "nit": digitos(f[5]), "dv": digitos(f[6]),
            "razon": texto(f[7]), "aplicacion": texto(f[9]), "correo": str(f[10] or "").strip(),
            "estado": texto(f[11]) or "ACTIVA", "obs": str(f[12] or "").strip(),
        })
    return out, sin_cuenta


def main():
    config = Path(sys.argv[1]) if len(sys.argv) > 1 else CONFIG_POR_DEFECTO
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else SALIDA_POR_DEFECTO
    avisos = []

    terceros, cuentas = leer_terceros(config / "BASE_TERCEROS.xlsx", avisos)
    cuentas_emp, sin_cuenta = leer_cuentas_empresa(config / "BASE_CUENTAS.xlsx", avisos)

    sql = ["-- Carga de terceros y cuentas de la empresa (A&F). GENERADO: no subir a git.",
           "-- Correr DESPUÉS de 20261016000000_ayf_terceros.sql.",
           "begin;", ""]

    sql.append("-- ------------------------------------------------ terceros")
    for lote in lotes(terceros):
        vals = ",\n  ".join(
            f"({q(t['tipo'])}, {q(t['doc'])}, {q(t['dv'])}, {q(t['nombre'])}, "
            f"{q(t['nombre_banco'])}, {q(t['nombres'])}, {q(t['apellidos'])})"
            for t in lote)
        sql.append("insert into public.terceros "
                   "(tipo_documento, numero_documento, dv, razon_social, nombre_banco, nombres, apellidos)\n"
                   f"values\n  {vals}\non conflict (numero_documento) do nothing;\n")

    sql.append("-- ------------------------------------- cuentas de terceros")
    for lote in lotes(cuentas):
        vals = ",\n  ".join(
            f"({q(c['doc'])}, {q(c['banco'])}, {q(c['tipo'])}, {q(c['cuenta'])}, {q(c['etiqueta'])}, "
            f"{q(c['estado'])}, {q(c['ver_en'])}, {q(c['ver_por'])}, {q(c['obs'])})"
            for c in lote)
        sql.append(
            "insert into public.terceros_cuentas\n"
            "  (tercero_id, banco_id, tipo_cuenta, numero_cuenta, etiqueta, estado,\n"
            "   verificada_en, verificada_por, observaciones)\n"
            "select t.id, b.id, v.tipo, v.cuenta, v.etiqueta, v.estado,\n"
            "       v.ver_en::date, v.ver_por, v.obs\n"
            f"from (values\n  {vals}\n"
            ") as v(doc, banco, tipo, cuenta, etiqueta, estado, ver_en, ver_por, obs)\n"
            "join public.terceros t on t.numero_documento = v.doc\n"
            "join public.bancos b on b.nombre = v.banco\n"
            "on conflict (tercero_id, banco_id, numero_cuenta) do nothing;\n")

    sql.append("-- ------------------------------------ cuentas de la empresa")
    if cuentas_emp:
        vals = ",\n  ".join(
            f"({q(c['apodo'])}, {q(c['dim'])}, {q(c['banco'])}, {q(c['tipo'])}, {q(c['cuenta'])}, "
            f"{q(c['nit'])}, {q(c['dv'])}, {q(c['razon'])}, {q(c['aplicacion'])}, {q(c['correo'])}, "
            f"{q(c['estado'])}, {q(c['obs'])})"
            for c in cuentas_emp)
        sql.append(
            "insert into public.cuentas_empresa\n"
            "  (apodo, diminutivo, banco_id, tipo_cuenta, numero_cuenta, nit, dv, razon_social,\n"
            "   aplicacion, correo_itau, estado, observaciones)\n"
            "select v.apodo, v.dim, b.id, v.tipo, v.cuenta, v.nit, v.dv, v.razon,\n"
            "       v.aplicacion, v.correo, v.estado, v.obs\n"
            f"from (values\n  {vals}\n"
            ") as v(apodo, dim, banco, tipo, cuenta, nit, dv, razon, aplicacion, correo, estado, obs)\n"
            "join public.bancos b on b.nombre = v.banco\n"
            "on conflict (apodo) do nothing;\n")
        sql.append(
            "-- Cruce con la tabla empresas por NIT (con o sin dígito de verificación).\n"
            "update public.cuentas_empresa ce\n"
            "   set empresa_id = e.id\n"
            "  from public.empresas e\n"
            " where ce.empresa_id is null\n"
            "   and regexp_replace(e.nit, '[^0-9]', '', 'g') in (ce.nit, ce.nit || coalesce(ce.dv, ''));\n")

    sql.append("commit;\n")
    sql.append("-- ------------------------------------------------ revisión")
    sql.append(
        "select 'terceros' as tabla, count(*) as filas from public.terceros\n"
        "union all select 'terceros_cuentas', count(*) from public.terceros_cuentas\n"
        "union all select 'cuentas_empresa', count(*) from public.cuentas_empresa;\n")
    sql.append(
        "-- Cuentas de la empresa que NO se pudieron cruzar con una empresa (revisar NIT):\n"
        "select apodo, razon_social, nit from public.cuentas_empresa where empresa_id is null;\n")
    sql.append(
        "-- Empresas sin consolidado asignado (asignar en Control administrativo > Empresas):\n"
        "select razon_social, nit from public.empresas where consolidado_id is null order by razon_social;\n")

    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text("\n".join(sql), encoding="utf-8")

    print(f"terceros: {len(terceros)} · cuentas de terceros: {len(cuentas)} · "
          f"cuentas de la empresa: {len(cuentas_emp)}")
    if sin_cuenta:
        print(f"empresas sin cuenta abierta (no se cargan): {', '.join(sin_cuenta)}")
    for a in avisos:
        print("AVISO:", a)
    print(f"escrito: {salida}")


if __name__ == "__main__":
    main()
