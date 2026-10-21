"""Genera el SQL con los datos iniciales de la liquidación de pagos (A&F,
migración 20261203000000_liquidacion_pagos.sql) desde:

  * la extracción de RUT (Extraccion_RUT.xlsx, hoja 01_EXTRACCION_RUT), y
  * el cuadro de pagos ya liquidados por Contabilidad (Pagos_Sem_31-42.xlsx,
    hoja Consolidado General).

Uso (desde app-demo/):
    python scripts/generar_datos_liquidacion.py RUT.xlsx PAGOS.xlsx [salida.sql]

Por defecto escribe referencias/seed_liquidacion_ayf.sql (carpeta en
.gitignore: trae documentos y nombres de terceros, NO debe subirse a git).

El SQL se corre a mano en el SQL Editor de Supabase, DESPUÉS de la migración.
Es repetible: no duplica ni pisa lo que Contabilidad ya haya decidido.
Carga:
  1. rut_extracciones (solo datos tributarios: sin correo, teléfono ni
     dirección) y el perfil tributario de los terceros que ya existen.
  2. Qué retiene cada empresa que paga (inferido del historial, por confirmar).
  3. La memoria por tercero: concepto y tarifa de retefuente, tarifa de ICA por
     municipio y ReteIVA, tal como los liquidó Contabilidad.
"""

import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

SALIDA_POR_DEFECTO = Path(__file__).resolve().parent.parent / "referencias" / "seed_liquidacion_ayf.sql"
LOTE = 300

# Código del proyecto -> municipio de ICA (igual que la migración).
MUNICIPIO_POR_CODIGO = {
    "CTG": ("BOLIVAR", "CARTAGENA"),
    "BAQ": ("ATLANTICO", "BARRANQUILLA"),
    "ARA": ("ARAUCA", "ARAUCA"),
    "AQT": ("ARAUCA", "ARAUQUITA"),
    "SAR": ("ARAUCA", "SARAVENA"),
    "BOG": ("BOGOTA", "BOGOTA"),
}

# Tarifas de ICA por municipio y actividad (las mismas de la migración), para
# saber qué actividad corresponde a la tarifa que usó Contabilidad.
ACTIVIDADES_ICA = {
    ("BOLIVAR", "CARTAGENA"): {"construccion": 0.00856, "compras": 0.00749, "hotel": 0.00642, "consultoria": 0.00321},
    ("ATLANTICO", "BARRANQUILLA"): {"construccion": 0.0125, "compras": 0.01, "transporte": 0.008},
    ("ARAUCA", "ARAUCA"): {"construccion": 0.01, "compras": 0.006, "transporte": 0.008, "hotel": 0.0085},
    ("ARAUCA", "ARAUQUITA"): {"construccion": 0.01, "compras": 0.005, "hotel": 0.008},
    ("ARAUCA", "SARAVENA"): {"construccion": 0.007, "compras": 0.005, "transporte": 0.006, "hotel": 0.006},
    ("BOGOTA", "BOGOTA"): {"construccion": 0.0076, "compras": 0.0069, "transporte": 0.00414, "consultoria": 0.00866,
                           "servicios_profesionales": 0.00766, "mantenimiento": 0.00966, "hotel": 0.0138},
}

# Clasificación del concepto escrito en el cuadro (solo para separar tarifas
# que comparten valor: 3,5 % compras / arriendo de inmueble / hotel...).
REGLAS_TEXTO = [
    ("no_aplica", r"N[OÓ]MINA|SEGURIDAD\s+SOCIAL|PLANILLA|PRESTACIONES|LIQUIDACI[OÓ]N\b|\bRT\b|AUTO ?RENTA|DIAN|IMPUESTO|"
                  r"ACUEDUCTO|ALCANTARILLADO|ENERG[IÍ]A|\bLUZ\b|GASES|INTERNET|TELEFON|CELULAR|P[OÓ]LIZA|SOAT|SEGURO|"
                  r"CAJA MENOR|REEMBOLSO|PEAJE|PR[EÉ]STAMO"),
    ("combustible", r"COMBUSTIBLE|DIESEL|GASOLINA|ACPM"),
    ("honorarios", r"SERVICIOS? PROFESIONAL|HONORARIO|ASESOR[IÍ]A|CONSULTOR|TOPOGRAF|AVAL[UÚ]O|DICTAMEN|RETIE|INTERVENTOR|DISE[NÑ]O"),
    ("transporte", r"TRANSPORTE|FLETE|ACARREO|MOTOCARRO|VOLQUETA|RECOLECCI"),
    ("arriendo_inmueble", r"CANON|ARRIENDO|ARRENDAMIENTO"),
    ("arriendo_mueble", r"ALQUILER"),
    ("hotel", r"HOSPEDAJE|HOTEL|ALOJAMIENTO|ALIMENTACI[OÓ]N|RESTAURANTE"),
    ("construccion", r"MANO DE OBRA|CORTE (DE|FINAL)|\bOBRA\b|OBRAS|ACTA|INSTALACI[OÓ]N"),
    ("servicios", r"SERVICIO|MANTENIMIENTO|ARREGLO|REPARACI|EX[AÁ]MEN"),
]
REGLAS_TEXTO = [(k, re.compile(v)) for k, v in REGLAS_TEXTO]


def clasificar(txt) -> str:
    t = str(txt or "").upper()
    for k, rx in REGLAS_TEXTO:
        if rx.search(t):
            return k
    return "compras"


def concepto_de_tarifa(tarifa: float, cat: str):
    """Concepto del catálogo que corresponde a la tarifa que usó Contabilidad."""
    if tarifa in (0.025, 0.0125):
        return "compras"
    if tarifa == 0.035:
        return {"arriendo_inmueble": "arrendamiento_inmuebles", "hotel": "hoteles_restaurantes",
                "transporte": "transporte_pasajeros"}.get(cat, "compras")
    if tarifa in (0.01, 0.005):
        return "transporte_carga"
    if tarifa == 0.02:
        return "construccion"
    if tarifa == 0.04:
        return "arrendamiento_muebles" if cat == "arriendo_mueble" else "servicios"
    if tarifa in (0.06, 0.03):
        return "servicios"
    if tarifa in (0.1, 0.11, 0.05, 0.055):
        return "honorarios"
    if tarifa in (0.001, 0.0005):
        return "combustibles"
    return None


def texto(v) -> str:
    s = str(v if v is not None else "").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def digitos(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip().split("-")[0]
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    return re.sub(r"\D", "", s)


def documento_pago(v) -> str:
    """Documento del tercero en el cuadro: sin puntos ni DV (un NIT de 10 dígitos
    que empieza por 8 o 9 trae el DV pegado)."""
    d = digitos(v)
    if len(d) == 10 and d[0] in "89":
        d = d[:-1]
    return d


def numero(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            return None


def fecha_iso(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.year >= 2020 else ""
    s = str(v or "").strip()[:10]
    return s if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", s) else ""


def dv_dian(nit: str) -> str:
    pesos = [3, 7, 13, 17, 19, 23, 29, 37, 41, 43, 47, 53, 59, 67, 71]
    s = sum(int(c) * pesos[i] for i, c in enumerate(reversed(nit)) if i < len(pesos))
    r = s % 11
    return str(r if r in (0, 1) else 11 - r)


TILDES = str.maketrans("ÁÉÍÓÚÜÑáéíóúüñ", "AEIOUUNAEIOUUN")


def clave_empresa(nombre: str) -> str:
    """Igual que public._clave_nombre_empresa."""
    s = (nombre or "").translate(TILDES).upper()
    s = re.sub(r"[^A-Z0-9 ]", "", s)
    s = re.sub(r"\b(SAS|SA|LTDA|ZOMAC|ZESE|EN REORGANIZACION)\b", "", s)
    return s.replace(" ", "")


def municipio_de_proyecto(codigo: str):
    c = (codigo or "").upper()
    for k, m in MUNICIPIO_POR_CODIGO.items():
        if re.search(rf"(^|-){k}(-|$| )", c):
            return m
    return None


def sql_txt(v) -> str:
    if v is None or v == "":
        return "null"
    return "'" + str(v).replace("'", "''") + "'"


def sql_arr(lista) -> str:
    return "'{" + ",".join(lista) + "}'::text[]" if lista else "'{}'::text[]"


# ----------------------------------------------------------------------- RUT
def leer_ruts(ruta: Path):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    encabezado = None
    for f in filas:
        if f and "archivo_origen" in [str(x) for x in f]:
            encabezado = [str(x) if x is not None else "" for x in f]
            break
    if encabezado is None:
        raise SystemExit("No encontré el encabezado de la extracción de RUT (archivo_origen...).")
    i = {n: k for k, n in enumerate(encabezado)}
    por_doc = {}
    for f in filas:
        if not f or all(x is None for x in f):
            continue
        g = lambda n: f[i[n]] if n in i and i[n] < len(f) else None
        nit = digitos(g("nit_extraido"))
        if not nit or len(nit) < 5:
            continue
        tipo_doc = texto(g("tipo_documento_extraido")).upper()
        if tipo_doc in ("NIT", "PERSONA JURÍDICA", "PERSONA JURIDICA", "31"):
            tipo = "juridica"
        elif tipo_doc:
            tipo = "natural"
        else:
            tipo = "juridica" if len(nit) == 9 and nit[0] in "89" else "natural"
        dv = digitos(g("dv_extraido"))[:1]
        codigos = sorted({c.zfill(2) for c in re.findall(r"\d+", str(g("codigos_casilla_53") or ""))})
        nombre = texto(g("razon_social")) or " ".join(
            texto(g(n)) for n in ("primer_nombre", "otros_nombres", "primer_apellido", "segundo_apellido") if texto(g(n)))
        fila = {
            "nit": nit, "dv": dv or None, "dv_cuadra": (dv == dv_dian(nit)) if dv else None, "tipo": tipo,
            "nombre": nombre.upper() or None, "depto": texto(g("departamento_rut")) or None,
            "municipio": texto(g("municipio_rut")) or None,
            "ciiu1": digitos(g("ciiu_principal"))[:4] or None, "ciiu2": digitos(g("ciiu_secundaria"))[:4] or None,
            "resp": codigos, "fecha": fecha_iso(g("fecha_actualizacion_rut")) or fecha_iso(g("fecha_expedicion_rut")),
            "calidad": texto(g("calidad_lectura")).upper() or None, "archivo": texto(g("archivo_origen")) or None,
        }
        previa = por_doc.get(nit)
        if previa is None or (fila["fecha"] or "") >= (previa["fecha"] or ""):
            por_doc[nit] = fila
    return por_doc


# -------------------------------------------------------------------- pagos
def leer_pagos(ruta: Path):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    encabezado = [texto(x).upper() for x in next(filas)]
    i = {n: k for k, n in enumerate(encabezado) if n}
    pagos = []
    for f in filas:
        if not f or all(x is None for x in f):
            continue
        g = lambda n: f[i[n]] if n in i and i[n] < len(f) else None
        pagos.append({
            "empresa": texto(g("GRUPO EMPRESARIAL")).upper(),
            "doc": documento_pago(g("NIT DE LA EMPRESA QUE FACTURA")),
            "concepto": texto(g("CONCEPTO DE LA FACTURA")),
            "proyecto": texto(g("CODIGO DE PROYECTO")).upper(),
            "base": numero(g("BASE")),
            "iva": numero(g("IVA")),
            "rete": numero(g("RETE")) or 0,
            "ica": numero(g("RETENCION DE ICA")) or 0,
            "reteiva": numero(g("RETENCION DE IVA")) or 0,
            "fecha": fecha_iso(g("FECHA DE PAGO")),
        })
    return pagos


def preferencias(pagos):
    ret = defaultdict(list)    # doc -> [(tarifa, cat, fecha)]
    ica = defaultdict(list)    # (doc, municipio) -> [(tarifa, fecha)]
    iva = defaultdict(list)    # doc -> [fecha]
    for p in pagos:
        if not p["doc"] or not p["base"] or p["base"] <= 0:
            continue
        if re.search(r"REEMBOLSO|CAJA MENOR|LEGALIZACI", p["concepto"].upper()):
            continue  # varias facturas en un pago: la tarifa sale mezclada
        if p["rete"] > 0:
            ret[p["doc"]].append((round(p["rete"] / p["base"], 4), clasificar(p["concepto"]), p["fecha"]))
        mun = municipio_de_proyecto(p["proyecto"])
        if p["ica"] > 0 and mun:
            ica[(p["doc"], mun)].append((round(p["ica"] / p["base"], 5), p["fecha"]))
        if p["reteiva"] > 0:
            iva[p["doc"]].append(p["fecha"])

    salida = []
    for doc, lista in ret.items():
        tarifa, veces = Counter(t for t, _, _ in lista).most_common(1)[0]
        cat = Counter(c for t, c, _ in lista if t == tarifa).most_common(1)[0][0]
        concepto = concepto_de_tarifa(tarifa, cat)
        if concepto is None:
            continue
        ultima = max((f for t, _, f in lista if t == tarifa and f), default="")
        salida.append((doc, "retefuente", None, concepto, None, tarifa, veces, ultima))
    for (doc, mun), lista in ica.items():
        tarifa, veces = Counter(t for t, _ in lista).most_common(1)[0]
        actividad = next((a for a, v in ACTIVIDADES_ICA.get(mun, {}).items() if abs(v - tarifa) < 1e-6), None)
        ultima = max((f for t, f in lista if t == tarifa and f), default="")
        salida.append((doc, "reteica", mun, None, actividad, tarifa, veces, ultima))
    for doc, fechas in iva.items():
        salida.append((doc, "reteiva", None, None, None, 0.15, len(fechas), max((f for f in fechas if f), default="")))
    return salida


def empresas(pagos, ruts):
    datos = defaultdict(lambda: {"nombres": set(), "n": 0, "con_rete": 0, "evid_rete": 0,
                                 "con_reteiva": 0, "evid_reteiva": 0, "ica": set()})
    for p in pagos:
        if not p["empresa"]:
            continue
        clave = clave_empresa(p["empresa"])
        e = datos[clave]
        e["nombres"].add(p["empresa"])
        e["n"] += 1
        cat = clasificar(p["concepto"])
        if p["rete"] > 0:
            e["con_rete"] += 1
        if (p["base"] or 0) >= 524000 and cat != "no_aplica" and "ANTICIPO" not in p["concepto"].upper():
            e["evid_rete"] += 1
        r = ruts.get(p["doc"])
        if r and "47" in r["resp"] and (p["iva"] or 0) > 0:
            e["evid_reteiva"] += 1
        if p["reteiva"] > 0:
            e["con_reteiva"] += 1
        mun = municipio_de_proyecto(p["proyecto"])
        if p["ica"] > 0 and mun:
            e["ica"].add(mun)
    return datos


# -------------------------------------------------------------------- SQL
def lotes(valores, tam=LOTE):
    for k in range(0, len(valores), tam):
        yield valores[k:k + tam]


def generar(ruts, prefs, emps) -> str:
    out = [
        "-- Datos iniciales de la liquidación de pagos (A&F). Generado por",
        "-- scripts/generar_datos_liquidacion.py el " + datetime.now().strftime("%Y-%m-%d %H:%M") + ".",
        "-- Correr DESPUÉS de 20261203000000_liquidacion_pagos.sql. Repetible.",
        "-- NO subir a git: trae documentos y nombres de terceros.",
        "",
        "begin;",
        "",
        "-- 1. Extracción de RUT (la más reciente por documento)",
    ]
    filas = []
    for r in ruts.values():
        dv_cuadra = "null" if r["dv_cuadra"] is None else ("true" if r["dv_cuadra"] else "false")
        filas.append(
            f"({sql_txt(r['nit'])},{sql_txt(r['dv'])},{dv_cuadra},{sql_txt(r['tipo'])},{sql_txt(r['nombre'])},"
            f"{sql_txt(r['depto'])},{sql_txt(r['municipio'])},{sql_txt(r['ciiu1'])},{sql_txt(r['ciiu2'])},"
            f"{sql_arr(r['resp'])},{sql_txt(r['fecha']) + '::date' if r['fecha'] else 'null'},"
            f"{sql_txt(r['calidad'])},{sql_txt(r['archivo'])})")
    for lote in lotes(filas):
        out.append("insert into public.rut_extracciones (numero_documento, dv, dv_cuadra, tipo_persona, nombre, "
                   "departamento, municipio, ciiu_principal, ciiu_secundaria, responsabilidades, fecha_actualizacion, "
                   "calidad_lectura, archivo_origen) values")
        out.append(",\n".join(lote))
        out.append("on conflict (numero_documento) do update set dv = excluded.dv, dv_cuadra = excluded.dv_cuadra, "
                   "tipo_persona = excluded.tipo_persona, nombre = excluded.nombre, departamento = excluded.departamento, "
                   "municipio = excluded.municipio, ciiu_principal = excluded.ciiu_principal, "
                   "ciiu_secundaria = excluded.ciiu_secundaria, responsabilidades = excluded.responsabilidades, "
                   "fecha_actualizacion = excluded.fecha_actualizacion, calidad_lectura = excluded.calidad_lectura, "
                   "archivo_origen = excluded.archivo_origen, cargado_en = now()")
        out.append("where coalesce(excluded.fecha_actualizacion, '1900-01-01') >= "
                   "coalesce(public.rut_extracciones.fecha_actualizacion, '1900-01-01');")
        out.append("")

    out += [
        "-- Perfil tributario de los terceros que ya existen (los nuevos lo reciben",
        "-- solos al crearse). No pisa perfiles existentes.",
        "do $$ begin",
        "  perform public._perfil_tributario_desde_rut(t.id)",
        "     from public.terceros t",
        "     join public.rut_extracciones r on r.numero_documento = t.numero_documento;",
        "end $$;",
        "",
        "-- 2. Qué retiene cada empresa (inferido de las semanas 31 a 42; por confirmar)",
    ]
    filas_emp, filas_ica = [], []
    for clave, e in sorted(emps.items()):
        if not clave:
            continue
        rf = e["con_rete"] > 0 or e["evid_rete"] < 5
        rv = e["con_reteiva"] > 0 or e["evid_reteiva"] == 0
        obs = (f"Inferido de las semanas 31-42 ({', '.join(sorted(e['nombres']))}): {e['n']} pagos, "
               f"{e['con_rete']} con retefuente, {e['con_reteiva']} con ReteIVA.")
        filas_emp.append(f"({sql_txt(clave)},{'true' if rf else 'false'},{'true' if rv else 'false'},{sql_txt(obs)})")
        for depto, mun in sorted(e["ica"]):
            filas_ica.append(f"({sql_txt(clave)},{sql_txt(depto)},{sql_txt(mun)})")
    out.append("insert into public.empresas_retenciones (empresa_id, agente_retefuente, agente_reteiva, confirmado, observaciones)")
    out.append("select e.id, v.rf, v.rv, false, v.obs")
    out.append("  from (values")
    out.append(",\n".join(filas_emp))
    out.append(") as v(clave, rf, rv, obs)")
    out.append("  join public.empresas e on public._clave_nombre_empresa(e.razon_social) = v.clave")
    out.append("on conflict (empresa_id) do nothing;")
    out.append("")
    if filas_ica:
        out.append("insert into public.empresas_reteica_municipios (empresa_id, municipio_id, confirmado)")
        out.append("select e.id, m.id, false")
        out.append("  from (values")
        out.append(",\n".join(filas_ica))
        out.append(") as v(clave, departamento, municipio)")
        out.append("  join public.empresas e on public._clave_nombre_empresa(e.razon_social) = v.clave")
        out.append("  join public.municipios m on m.departamento = v.departamento and m.nombre = v.municipio")
        out.append("on conflict (empresa_id, municipio_id) do nothing;")
        out.append("")
    out.append("-- Empresas del cuadro que no están en la app (o tienen otro nombre): revisar.")
    out.append("select v.clave as empresa_del_cuadro_sin_cuadrar from (values")
    out.append(",\n".join(f"({sql_txt(c)})" for c in sorted(emps) if c))
    out.append(") as v(clave)")
    out.append("where not exists (select 1 from public.empresas e where public._clave_nombre_empresa(e.razon_social) = v.clave);")
    out.append("")

    out.append("-- 3. Memoria por tercero (lo que Contabilidad liquidó). No pisa lo que")
    out.append("-- Contabilidad ya haya decidido en la app.")
    filas_pref = []
    for doc, tipo, mun, concepto, actividad, tarifa, veces, ultima in prefs:
        depto, nom = mun if mun else (None, None)
        filas_pref.append(
            f"({sql_txt(doc)},{sql_txt(tipo)},{sql_txt(depto)},{sql_txt(nom)},{sql_txt(concepto)},{sql_txt(actividad)},"
            f"{tarifa},{veces},{sql_txt(ultima) + '::date' if ultima else 'null'})")
    for lote in lotes(filas_pref):
        out.append("insert into public.liquidacion_preferencias (numero_documento, tipo, municipio_id, concepto_clave, "
                   "actividad_ica, tarifa, veces, origen, ultima_vez)")
        out.append("select v.doc, v.tipo, m.id, v.concepto, v.actividad, v.tarifa, v.veces, 'historial', "
                   "coalesce(v.ultima::timestamptz, now())")
        out.append("  from (values")
        out.append(",\n".join(lote))
        out.append(") as v(doc, tipo, departamento, municipio, concepto, actividad, tarifa, veces, ultima)")
        out.append("  left join public.municipios m on m.departamento = v.departamento and m.nombre = v.municipio")
        out.append(" where v.tipo <> 'reteica' or m.id is not null")
        out.append("on conflict (numero_documento, tipo, coalesce(municipio_id, '00000000-0000-0000-0000-000000000000'::uuid)) do nothing;")
        out.append("")

    out += [
        "commit;",
        "",
        "-- Resumen",
        "select (select count(*) from public.rut_extracciones) as rut_cargados,",
        "       (select count(*) from public.terceros_tributario) as terceros_con_perfil,",
        "       (select count(*) from public.terceros t where not exists",
        "          (select 1 from public.terceros_tributario x where x.tercero_id = t.id)) as terceros_sin_perfil,",
        "       (select count(*) from public.empresas_retenciones) as empresas_con_retenciones,",
        "       (select count(*) from public.empresas_reteica_municipios) as empresas_municipios_ica,",
        "       (select count(*) from public.liquidacion_preferencias) as preferencias,",
        "       (select count(*) from public.proyectos where municipio_ica_id is null) as proyectos_sin_municipio_ica;",
        "",
    ]
    return "\n".join(out)


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    ruts = leer_ruts(Path(sys.argv[1]))
    pagos = leer_pagos(Path(sys.argv[2]))
    salida = Path(sys.argv[3]) if len(sys.argv) > 3 else SALIDA_POR_DEFECTO
    prefs = preferencias(pagos)
    emps = empresas(pagos, ruts)
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(generar(ruts, prefs, emps), encoding="utf-8")
    tipos = Counter(p[1] for p in prefs)
    print(f"RUT: {len(ruts)} documentos · pagos leídos: {len(pagos)} · empresas: {len(emps)}")
    print(f"Preferencias: {dict(tipos)}")
    print(f"Escrito en {salida}")


if __name__ == "__main__":
    main()
